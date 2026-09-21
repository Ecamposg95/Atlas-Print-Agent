"""Lectura de catálogos .xlsx y .csv con mapeo de columnas por alias."""

from __future__ import annotations

import csv
import unicodedata
from pathlib import Path
from typing import Iterable

from .model import Product, parse_price, parse_stock


class CatalogError(Exception):
    """Error legible al leer o interpretar el catálogo."""


# Primer alias presente gana. Encabezados normalizados con normalize_header.
ALIASES: dict[str, tuple[str, ...]] = {
    "sku": ("sku",),
    "name": ("nombreventa", "nombre", "producto"),
    "brand": ("marca",),
    "barcode": ("codigobarras", "barcode", "codigo"),
    "price": ("preciotexto", "preciobase", "precio"),
    "stock": ("stock", "cantidad", "existencia"),
    "color": ("color",),
    "size": ("talla",),
}
REQUIRED = ("sku", "name")


def normalize_header(header) -> str:
    text = unicodedata.normalize("NFKD", str(header or "")).encode("ascii", "ignore").decode()
    return "".join(ch for ch in text.lower() if ch.isalnum())


def map_columns(headers: list) -> dict[str, int]:
    positions: dict[str, int] = {}
    for idx, header in enumerate(headers):
        key = normalize_header(header)
        if key and key not in positions:
            positions[key] = idx
    mapping: dict[str, int] = {}
    for field, aliases in ALIASES.items():
        for alias in aliases:
            if alias in positions:
                mapping[field] = positions[alias]
                break
    missing = [f for f in REQUIRED if f not in mapping]
    if missing:
        found = ", ".join(str(h) for h in headers if h)
        raise CatalogError(
            f"Faltan columnas obligatorias: {', '.join(missing)}. Encabezados encontrados: {found}"
        )
    return mapping


def _cell(row, idx: int | None) -> str:
    if idx is None or idx >= len(row):
        return ""
    value = row[idx]
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def rows_to_products(headers: list, rows: Iterable) -> list[Product]:
    mapping = map_columns(headers)
    products: list[Product] = []
    for row in rows:
        get = lambda field: _cell(row, mapping.get(field))  # noqa: E731
        sku, name = get("sku"), get("name")
        if not sku and not name:
            continue
        price, price_text = parse_price(get("price"))
        products.append(Product(
            sku=sku, name=name, brand=get("brand"), barcode=get("barcode"),
            price=price, price_text=price_text, stock=parse_stock(get("stock")),
            color=get("color"), size=get("size"),
        ))
    return products


def sheet_names(path) -> list[str]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _read_xlsx(path: Path, sheet: str | None) -> list[Product]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet is not None:
            if sheet not in wb.sheetnames:
                raise CatalogError(f"La hoja {sheet!r} no existe. Hojas: {', '.join(wb.sheetnames)}")
            ws = wb[sheet]
        else:
            ws = wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        headers = next(rows, None)
        if headers is None:
            raise CatalogError(f"La hoja {ws.title!r} está vacía")
        return rows_to_products(list(headers), rows)
    finally:
        wb.close()


def _read_csv(path: Path) -> list[Product]:
    with open(path, "r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.reader(fh)
        headers = next(reader, None)
        if headers is None:
            raise CatalogError(f"El archivo {path.name} está vacío")
        return rows_to_products(headers, reader)


def read_catalog(path, sheet: str | None = None) -> list[Product]:
    p = Path(path)
    if not p.exists():
        raise CatalogError(f"No existe el archivo: {p}")
    suffix = p.suffix.lower()
    if suffix == ".csv":
        return _read_csv(p)
    if suffix in (".xlsx", ".xlsm"):
        return _read_xlsx(p, sheet)
    raise CatalogError(f"Formato no soportado: {p.suffix}. Usa .xlsx o .csv")


def select(products: list[Product], skus: list[str] | None = None, search: str | None = None) -> list[Product]:
    wanted = {s.strip().upper() for s in skus if s.strip()} if skus else None
    term = (search or "").strip().lower()
    out: list[Product] = []
    for p in products:
        if wanted is not None and p.sku.upper() not in wanted:
            continue
        if term and term not in " ".join((p.sku, p.barcode, p.brand, p.name)).lower():
            continue
        out.append(p)
    return out
